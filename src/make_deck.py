# ============================================================
# make_deck.py
# Purpose : Build the presentation as a .pptx in the course's KiTH
#           template style (imports into Canva, PowerPoint, Google
#           Slides). All numbers are read from results/, so re-run it
#           after any new results.
#
# One-time setup:  venv\Scripts\pip.exe install python-pptx
# How to run:      venv\Scripts\python.exe src\make_deck.py
# Output:          slides/khmer_sign_language_presentation.pptx
# ============================================================

import json
import os
import sys

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from lxml import etree
from matplotlib.colors import to_rgb
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import make_slide_art as art  # noqa: E402

RESULTS = "results"
ART = os.path.join("slides", "art")
CANVA = os.path.join(ART, "canva")
CHARTS = os.path.join(CANVA, "charts")
LOGO_RAW = os.path.join(ART, "kith_logo_raw.png")
LOGO = os.path.join(ART, "kith_logo.png")
PANEL = os.path.join(ART, "title_panel.png")
OUT = os.path.join("slides", "khmer_sign_language_presentation.pptx")
FONT, FOOT_FONT = "Arial", "Times New Roman"


def rgb(hex_):
    h = hex_.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


# Colours sampled from the template deck
GREEN, RULE, BAND = rgb("#0A593C"), rgb("#00553A"), rgb("#084B32")
BLACK, GRAY, LIGHT_LINE, WHITE = rgb("#111111"), rgb("#555555"), rgb("#DADFDC"), rgb("#FFFFFF")
ROW_ALT, GOOD_ROW = rgb("#F1F6F3"), rgb("#DDF0E6")
MODEL_COLOR = {"lstm": rgb("#C98600"), "gru": rgb("#0072B2"), "transformer": rgb("#009E73")}


# ---------------------------------------------------------------- data
def J(*parts):
    with open(os.path.join(RESULTS, *parts)) as f:
        return json.load(f)


models = ["gru", "transformer", "lstm"]
tuned = os.path.exists(os.path.join(RESULTS, "lstm_tuned", "metrics.json"))
rows_models = models + (["lstm_tuned"] if tuned else [])
meta = {m: J(m, "meta.json") for m in rows_models}
met = {m: J(m, "metrics.json") for m in rows_models}
hist = {m: J(m, "history.json") for m in models}
diag = J("diagnostics.json")
split = J("split_info.json")
hp = {m: J(m, "hparam_search.json") for m in ("gru", "lstm")}
errors = pd.read_csv(os.path.join(RESULTS, "lstm", "errors.csv"))
pct = lambda x: f"{x * 100:.0f}%"
pretty = lambda s: s.replace("dont_understand", "don't understand").replace("_", " ")
NAME = {"lstm": "LSTM", "gru": "GRU", "transformer": "Transformer", "lstm_tuned": "LSTM (tuned)"}
acc = {m: met[m]["test_accuracy"] for m in rows_models}
ci = {m: diag["models"][m]["test_accuracy_ci95"] for m in models}
n_test = split["n_test"]
n_err = {m: round((1 - acc[m]) * n_test) for m in rows_models}
t_best = int(np.argmax(hist["transformer"]["val_acc"])) + 1
lstm_train_acc = hist["lstm"]["train_acc"][-1]
guess = errors[errors.predicted_label == "understand"]
guess_conf = guess.predicted_confidence.mean()
secs = [meta[m]["training_time_seconds"] for m in models]
gru_res, lstm_res = hp["gru"]["results"], hp["lstm"]["results"]
gru_full = [r for r in gru_res if r["best_val_acc"] >= 0.999]
lstm_default = next(r for r in lstm_res if r["lr"] == 1e-3 and r["weight_decay"] == 0.0)
lstm_best = hp["lstm"]["best"]
slow_gru = [r["best_val_acc"] for r in gru_res if r["lr"] == 1e-4]


# ---------------------------------------------------------------- template assets
def make_logo():
    """Crop the KiTH logo (taken from the template deck) and make its white background transparent."""
    im = np.array(Image.open(LOGO_RAW).convert("RGB")).astype(float)
    m = im.min(axis=2)
    alpha = np.clip((250 - m) / 50, 0, 1) * 255
    rgba = np.dstack([im, alpha]).astype(np.uint8)
    img = Image.fromarray(rgba, "RGBA")
    img = img.crop(img.getchannel("A").point(lambda a: 255 if a > 8 else 0).getbbox())
    img.save(LOGO)


def make_title_panel():
    """Rounded picture for the title slide (green panel with the signer drawing)."""
    w, h = 4.66, 3.98
    fig = plt.figure(figsize=(w, h), dpi=300)
    ax = fig.add_axes([0, 0, 1, 1])
    ax.set_xlim(0, w); ax.set_ylim(0, h); ax.axis("off")
    X, Y = np.meshgrid(np.linspace(0, 1, 320), np.linspace(0, 1, 280))
    c1, c2, glow = np.array(to_rgb("#06402B")), np.array(to_rgb("#0F6B47")), np.array(to_rgb("#63D3A0"))
    t = (0.6 * X + 0.4 * (1 - Y))[..., None]
    img = c1 * (1 - t) + c2 * t
    img = img + 0.22 * glow * np.exp(-(((X - 0.5) ** 2) / 0.10 + ((Y - 0.5) ** 2) / 0.12))[..., None]
    ax.imshow(np.clip(img, 0, 1), extent=[0, w, 0, h], origin="lower", aspect="auto", zorder=0)
    style = dict(pose_line="#E8F5EE", pose_dot="#FFD166", hand_line="#9BEBC6", hand_dot="#9BEBC6",
                 edge="white", glow=True)
    s = 0.37
    art.draw_signer(ax, (w / 2 - 5.0 * s, h / 2 - 5.75 * s), s, style, lw=2.2, ms=48)
    fig.canvas.draw()
    pic = Image.fromarray(np.asarray(fig.canvas.buffer_rgba())).convert("RGBA")
    plt.close(fig)
    mask = Image.new("L", pic.size, 0)
    ImageDraw.Draw(mask).rounded_rectangle([0, 0, pic.size[0] - 1, pic.size[1] - 1], radius=110, fill=255)
    pic.putalpha(mask)
    pic.save(PANEL)


make_logo()
make_title_panel()

# ---------------------------------------------------------------- helpers
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def set_bullet(p, char="•", indent=0.3):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(indent))))
    pPr.set("indent", str(-int(Inches(indent))))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    etree.SubElement(pPr, qn("a:buChar")).set("char", char)


def text(slide, x, y, w, h, paras, size=16, color=BLACK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, bullets=False, after=12, spacing=1.08, font=FONT):
    """paras: list of str, or list of lists of (text, {size,bold,color}) runs."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(after)
        p.line_spacing = spacing
        runs = [(para, {})] if isinstance(para, str) else para
        for t, o in runs:
            r = p.add_run()
            r.text = t
            r.font.name = o.get("font", font)
            r.font.size = Pt(o.get("size", size))
            r.font.bold = o.get("bold", bold)
            r.font.color.rgb = o.get("color", color)
        if bullets:
            set_bullet(p)
    return box


def rect(slide, x, y, w, h, fill):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def picture(slide, path, bx, by, bw, bh, align="center"):
    iw, ih = Image.open(path).size
    scale = min(bw / iw, bh / ih)
    w, h = iw * scale, ih * scale
    x = bx + (bw - w) / 2 if align == "center" else bx
    y = by + (bh - h) / 2
    return slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))


def chrome(slide, rule_y=0.96):
    """Logo, green rule and footer band, as in the template deck."""
    lw, lh = Image.open(LOGO).size
    slide.shapes.add_picture(LOGO, Inches(0.68), Inches(0.2), height=Inches(0.58))
    rect(slide, 0.76, rule_y, 11.94, 0.04, RULE)
    rect(slide, 0, 7.0, 13.333, 0.5, BAND)
    text(slide, 0.23, 7.0, 6, 0.5, ["Student final project"], size=15, color=WHITE,
         anchor=MSO_ANCHOR.MIDDLE, font=FOOT_FONT, after=0)
    text(slide, 7.1, 7.0, 6.0, 0.5, ["Deep Learning"], size=15, color=WHITE, align=PP_ALIGN.RIGHT,
         anchor=MSO_ANCHOR.MIDDLE, font=FOOT_FONT, after=0)


def new_slide(title=None, subtitle=None):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = WHITE
    chrome(s)
    if title:
        text(s, 0.76, 1.22, 11.9, 0.6, [title], size=30, bold=True, color=GREEN, after=0)
    if subtitle:
        text(s, 0.78, 1.82, 11.9, 0.3, [subtitle], size=14, color=GREEN, after=0)
    return s


def notes(slide, body):
    slide.notes_slide.notes_text_frame.text = body


def heading(slide, x, y, w, label):
    text(slide, x, y, w, 0.35, [label], size=16, bold=True, color=BLACK, after=0)


def block(slide, x, y, w, head, body, head_color=BLACK, body_size=15):
    """Bold heading with a grey description, like the template's 'Models and Strategies' slide."""
    text(slide, x, y, w, 0.4, [head], size=19, bold=True, color=head_color, after=0)
    text(slide, x, y + 0.5, w, 1.3, body if isinstance(body, list) else [body], size=body_size, color=GRAY, after=4)


def callout(slide, cx, y, number, label, w=3.3):
    text(slide, cx - w / 2, y, w, 0.6, [number], size=30, bold=True, color=BLACK, align=PP_ALIGN.CENTER, after=0)
    text(slide, cx - w / 2, y + 0.65, w, 0.4, [label], size=13, color=GRAY, align=PP_ALIGN.CENTER, after=0)


def table(slide, x, y, widths, rows, font=15, row_h=0.5, highlight=None):
    shape = slide.shapes.add_table(len(rows), len(widths), Inches(x), Inches(y),
                                   Inches(sum(widths)), Inches(row_h * len(rows)))
    tbl = shape.table
    tbl.first_row = True
    tbl.horz_banding = False
    for j, w in enumerate(widths):
        tbl.columns[j].width = Inches(w)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            c = tbl.cell(i, j)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.margin_left = c.margin_right = Inches(0.12)
            c.fill.solid()
            c.fill.fore_color.rgb = GREEN if i == 0 else (GOOD_ROW if highlight == i else (ROW_ALT if i % 2 == 0 else WHITE))
            tf = c.text_frame
            tf.paragraphs[0].alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            r = tf.paragraphs[0].add_run()
            r.text = str(val)
            r.font.name, r.font.size = FONT, Pt(font)
            r.font.bold = i == 0 or j == 0
            r.font.color.rgb = WHITE if i == 0 else BLACK
    return tbl


# ================================================================ slides
# 1 — Title
s = prs.slides.add_slide(BLANK)
s.background.fill.solid()
s.background.fill.fore_color.rgb = WHITE
chrome(s, rule_y=0.91)
text(s, 0.8, 1.55, 6.7, 1.6, ["Khmer Sign Language", "Word Recognition"], size=40, bold=True, color=GREEN,
     after=0, spacing=1.0)
text(s, 0.82, 3.2, 6.6, 0.8, ["Comparing LSTM, GRU and Transformer on landmark sequences"], size=18, color=GREEN)
text(s, 0.82, 4.15, 6.6, 2.3, ["Presenter: [Your full name]", "Course: Deep Learning", "Lecturer: Mr. Soklong HIM",
                                "School: Kirirom Institute of Technology", "Date: [Presentation date]"],
     size=16, color=BLACK, after=3)
rect(s, 7.68, 1.95, 0.012, 3.2, rgb("#444444"))
s.shapes.add_picture(PANEL, Inches(7.92), Inches(1.66), width=Inches(4.66))
notes(s, "Say your name. One sentence: I built a system that recognises 11 Khmer sign words and "
         "compared three neural networks on it.")

# 2 — Table of contents
s = new_slide("Table of Contents")
toc = ["Problem and Approach", "Dataset", "Model Architectures", "Experimental Setup", "Results",
       "Learning Curves and Tuning", "Error Analysis", "Discussion", "Limitations", "Conclusion and Next Steps"]
for i, label in enumerate(toc):
    col, row = i // 5, i % 5
    x0 = 1.0 + col * 5.95
    y0 = 2.25 + row * 0.88
    text(s, x0 + 0.05, y0, 0.7, 0.8, [f"{i + 1:02d}"], size=15, bold=True, color=GRAY, anchor=MSO_ANCHOR.MIDDLE, after=0)
    text(s, x0 + 0.85, y0, 4.2, 0.8, [label], size=19, color=BLACK, anchor=MSO_ANCHOR.MIDDLE, after=0)
    rect(s, x0, y0 + 0.82, 4.95, 0.012, LIGHT_LINE)
notes(s, "Ten topics, about one minute each or less. You can delete this slide if you are short of time.")

# 3 — Problem and approach
s = new_slide("Problem and Approach", "From a webcam clip to one of 11 Khmer sign words")
picture(s, os.path.join(CANVA, "problem_illustration.png"), 0.76, 2.55, 6.4, 3.9, align="left")
heading(s, 7.7, 2.45, 5.0, "What the system does")
text(s, 7.7, 3.0, 5.0, 1.8, [
    "Reads a short webcam clip of a Khmer sign",
    "Input: 90 frames × 258 landmark values (about 3 s)",
    "Output: one of 11 words",
], size=16, bullets=True, after=10)
heading(s, 7.7, 4.95, 5.0, "The question")
text(s, 7.7, 5.5, 5.0, 1.3, ["Which model reads the signs best: LSTM, GRU or Transformer?"],
     size=16, bullets=True, after=10)
notes(s, "Explain input and output precisely. MediaPipe turns each frame into 258 numbers; it is a fixed, "
         "pretrained step shared by all three models. Only the sequence model is trained, from scratch, in PyTorch.")

# 4 — Dataset
s = new_slide("Dataset", "Self-collected landmarks, 11 signs, fixed split")
callout(s, 2.7, 2.35, "330", "clips (30 per sign)")
callout(s, 6.67, 2.35, "11", "sign classes")
callout(s, 10.6, 2.35, f"{split['n_train']} / {split['n_val']} / {split['n_test']}", "train / validation / test")
picture(s, os.path.join(CANVA, "sign_cards_clean.png"), 0.76, 3.55, 6.5, 3.3, align="left")
heading(s, 7.9, 3.65, 4.8, "Notes")
text(s, 7.9, 4.2, 4.8, 2.6, [
    "One signer, one webcam",
    "Fixed, stratified split (seed 42)",
    "A hand the camera misses is stored as zeros",
], size=16, bullets=True, after=10)
notes(s, "Self-collected, so I describe it honestly: one person. Balanced, 30 clips per sign. "
         "Teal cards are one-handed signs, purple are two-handed. The test set is never used for any decision. "
         "[Add one line on recording conditions: sittings, lighting, camera.]")

# 5 — Architectures
s = new_slide("Model Architectures", "Three networks compared on the same task")
picture(s, os.path.join(CHARTS, "fig_architectures.png"), 0.6, 2.2, 8.2, 4.7, align="left")
y = 2.45
for key, desc in [("lstm", "Gated memory"), ("gru", "Simpler gates"), ("transformer", "Self-attention, no recurrence")]:
    text(s, 9.3, y, 3.5, 0.4, [NAME[key]], size=19, bold=True, color=MODEL_COLOR[key], after=0)
    text(s, 9.3, y + 0.45, 3.5, 0.4, [desc], size=15, color=GRAY, after=0)
    y += 1.1
text(s, 9.3, y + 0.1, 3.5, 0.5, ["All trained from scratch"], size=16, bold=True, color=BLACK, after=0)
notes(s, "LSTM and GRU read the clip frame by frame and keep a memory. The Transformer lets every frame look at "
         "every other frame. GRU has about three quarters of the LSTM's parameters: 3 weight blocks instead of 4.")

# 6 — Setup
s = new_slide("Experimental Setup", "Same data, same metrics, same training recipe")
block(s, 0.9, 2.45, 5.6, "Data and metrics",
      [f"{split['n_train']} / {split['n_val']} / {split['n_test']} split, same for every model",
       "Accuracy (main), macro F1, confusion matrix"])
block(s, 0.9, 4.55, 5.6, "Training",
      ["Cross-entropy loss, Adam, learning rate 1e-3, batch size 16", "Early stopping (patience 25), dropout 0.2"])
block(s, 7.1, 2.45, 5.6, "Tuning",
      ["Learning rate {1e-3, 5e-4, 1e-4} × weight decay {0, 1e-4}", "Chosen on validation accuracy"])
block(s, 7.1, 4.55, 5.6, "Reproducibility and hardware",
      ["Seed 42, torch.save checkpoints",
       f"Colab Tesla T4, {min(secs):.0f}–{max(secs):.0f} s per model, [PyTorch version]"])
notes(s, "One sentence per block. The best epoch, chosen by validation accuracy, is what gets tested. "
         "Fill in the PyTorch version from the second cell of your Colab notebook.")

# 7 — Results
s = new_slide("Results", "Same 50 test clips for every model")
order = sorted(rows_models, key=lambda m: -acc[m])
n = len(order)
for i, m in enumerate(order):
    callout(s, 0.76 + (i + 0.5) * (11.94 / n), 2.3, pct(acc[m]), f"{NAME[m]} test accuracy", w=11.94 / n - 0.2)
trows = [["Approach", "Accuracy", "Macro F1", "Parameters", "Time"]]
for m in order:
    trows.append([NAME[m], pct(acc[m]), f"{met[m]['f1_macro']:.2f}", f"{meta[m]['trainable_params']:,}",
                  f"{meta[m]['training_time_seconds']:.1f} s"])
table(s, 0.9, 3.65, [2.2, 1.4, 1.3, 1.7, 1.1], trows, font=15, highlight=1)
heading(s, 9.1, 3.65, 3.6, "Reading the table")
text(s, 9.1, 4.15, 3.6, 2.7, [
    "Errors out of 50: " + ", ".join(f"{NAME[m]} {n_err[m]}" for m in order),
    f"95% intervals: GRU {pct(ci['gru'][0])}–{pct(ci['gru'][1])}, Transformer "
    f"{pct(ci['transformer'][0])}–{pct(ci['transformer'][1])}. Not separable with 50 clips",
], size=14, bullets=True, after=8)
notes(s, "Same test set, same metrics for every model. Be clear that the gap between GRU and Transformer is one clip.")

# 8 — Curves
s = new_slide("Learning Curves", "Training vs validation, per model")
picture(s, os.path.join(CHARTS, "fig_curves_train_val.png"), 0.6, 2.2, 8.7, 4.7, align="left")
y = 2.5
for key, line in [("lstm", f"Underfits: training accuracy only {pct(lstm_train_acc)}"),
                  ("gru", "Fits cleanly: no gap"),
                  ("transformer", f"Slight overfit: validation loss rises after epoch {t_best}")]:
    text(s, 9.7, y, 3.1, 0.4, [NAME[key]], size=19, bold=True, color=MODEL_COLOR[key], after=0)
    text(s, 9.7, y + 0.45, 3.1, 0.9, [line], size=15, color=GRAY, after=0)
    y += 1.4
notes(s, "Top row accuracy, bottom row loss on a log scale. In the Transformer's bottom panel the two lines separate: "
         "that is the overfitting signature from class. Early stopping kept the best epoch of each model.")

# 9 — Tuning
s = new_slide("Hyperparameter Tuning", "Learning rate × weight decay, chosen on validation")
picture(s, os.path.join(CHARTS, "fig_tuning.png"), 0.6, 2.3, 8.2, 4.4, align="left")
third = (f"Tuned LSTM test accuracy: {pct(acc['lstm_tuned'])}" if tuned
         else "The LSTM in the results uses the default rate; retraining it is future work")
heading(s, 9.2, 2.45, 3.6, "What we found")
text(s, 9.2, 3.0, 3.6, 3.9, [
    f"GRU: {len(gru_full)} of 6 settings reach 100%; learning rate 1e-4 is too slow "
    f"({pct(min(slow_gru))}–{pct(max(slow_gru))})",
    f"LSTM: a lower learning rate lifts validation accuracy from {pct(lstm_default['best_val_acc'])} "
    f"to {pct(lstm_best['best_val_acc'])}",
    third,
], size=14, bullets=True, after=10)
notes(s, "Choices were made on validation accuracy only. Weight decay made no difference for the GRU. "
         "This also tells us the LSTM's weak result is partly a settings problem.")

# 10 — Error analysis
s = new_slide("Error Analysis", "What each model got wrong, and why")
picture(s, os.path.join(CHARTS, "fig_lstm_errors.png"), 0.6, 2.4, 8.4, 4.2, align="left")
heading(s, 9.3, 2.45, 3.5, "What went wrong")
text(s, 9.3, 3.0, 3.5, 3.9, [
    "GRU: 0 errors",
    "Transformer: 1 error (understand read as don't understand)",
    f"LSTM: {n_err['lstm']} errors; {len(guess)} were “understand” at about {guess_conf:.2f} confidence "
    f"(chance is 0.09)",
    "All from one-handed signs: it learned which hands are present, not the movement",
], size=14, bullets=True, after=9)
notes(s, "The LSTM got sorry, thank you and understand fully right and collapsed the one-handed signs into one answer. "
         "Diagnosis: underfitting, not overfitting. I do not claim why the Transformer confused that one pair.")

# 11 — Discussion
s = new_slide("Why GRU Won, and a Caution", "Explaining the results with course concepts")
heading(s, 0.9, 2.45, 6.0, "Why")
text(s, 0.9, 3.0, 6.0, 3.8, [
    "GRU beats LSTM: little data (230 clips) and 90 steps; fewer parameters, easier to optimise",
    "Transformer is close: attention reaches every frame; the larger model overfits slightly",
    "But the dataset is easy: a model with no sense of order also scores 100%",
], size=16, bullets=True, after=12)
heading(s, 7.6, 2.45, 5.0, "Frame-order test")
dm = diag["models"]
drows = [["", "Normal order", "Frames shuffled"]]
for m in models:
    drows.append([NAME[m], pct(dm[m]["test_accuracy"]), pct(dm[m]["shuffled_frames_mean"])])
drows.append(["No-order baseline", pct(diag["baselines"]["logreg_clip_mean_no_order"]), "n/a"])
table(s, 7.6, 3.0, [2.2, 1.4, 1.5], drows, font=14, row_h=0.48)
text(s, 7.6, 5.6, 5.1, 1.2, ["Shuffling the frames barely hurts, so the models rely on hand shape and position "
                              "more than on movement."], size=14, color=GRAY)
notes(s, "This is the most important slide for honesty. A logistic regression on the clip average, with no concept "
         "of time, scores 100% on the test set. So the scores do not prove the models learned movement. "
         "A harder test: signs that differ only in movement, or a second signer. "
         "Course terms: capacity, data size, regularisation, inductive bias.")

# 12 — Limitations
s = new_slide("Limitations", "Honest limits of this study")
heading(s, 0.9, 2.45, 5.8, "Data and evaluation")
text(s, 0.9, 3.0, 5.8, 3.8, [
    "One signer, one room: no evidence for other people, cameras or lighting",
    "Small data: 30 clips per sign; one test clip is worth 2 points",
    "Easy dataset: frame order is barely used",
], size=16, bullets=True, after=12)
heading(s, 7.1, 2.45, 5.6, "Method")
text(s, 7.1, 3.0, 5.6, 3.8, [
    "Raw coordinates, and a missed hand becomes zeros: possible shortcuts",
    "No augmentation",
    "Live webcam accuracy not measured",
], size=16, bullets=True, after=12)
notes(s, "Read these straight; the rubric rewards limitations clearly identified. "
         "[If you ran a live test with predict.py, replace the last bullet with your tally.]")

# 13 — Conclusion
s = new_slide("Conclusion and Next Steps", "Key takeaways and prioritised future work")
heading(s, 0.9, 2.45, 5.8, "Conclusion")
text(s, 0.9, 3.0, 5.8, 3.8, [
    f"GRU best ({pct(acc['gru'])}), Transformer {pct(acc['transformer'])}, LSTM {pct(acc['lstm'])}",
    "Same data, split, seed and metrics for all three",
    "Honest finding: an easy dataset, order barely used",
], size=16, bullets=True, after=12)
heading(s, 7.1, 2.45, 5.6, "Next steps")
text(s, 7.1, 3.0, 5.6, 3.8, [
    "Record a second signer",
    "Add signs that differ only in movement",
    "Augmentation and body-relative coordinates",
    "Retrain the tuned LSTM; live accuracy test",
], size=16, bullets=True, after=12)
notes(s, "Three takeaways, then the next steps in priority order, then questions.")

# 14 — Thank you
s = new_slide()
text(s, 0.76, 3.0, 11.8, 1.2, ["Thank you"], size=48, bold=True, color=GREEN, align=PP_ALIGN.CENTER, after=0)
text(s, 0.76, 4.35, 11.8, 0.5, ["Questions?   ·   github.com/rinrin168/khmersignlanguage"], size=16, color=GRAY,
     align=PP_ALIGN.CENTER, after=0)

# Backup A1 — per sign
s = new_slide("Backup: Per-Sign Results", "For questions only")
tp, lp = met["transformer"]["per_class"], met["lstm"]["per_class"]
t_low = [f"{pretty(k)} {pct(v['recall'])}" for k, v in tp.items() if v["recall"] < 1]
t_prec = [f"{pretty(k)} {pct(v['precision'])}" for k, v in tp.items() if v["precision"] < 1]
groups = {}
for k, v in lp.items():
    groups.setdefault(round(v["recall"], 2), []).append(pretty(k))
heading(s, 0.9, 2.45, 5.8, "Transformer")
text(s, 0.9, 3.0, 5.8, 3.5, ["Recall 100% for every sign except: " + (", ".join(t_low) or "none"),
                              "Precision below 100%: " + (", ".join(t_prec) or "none")],
     size=15, bullets=True, after=10)
heading(s, 7.1, 2.45, 5.6, "LSTM (recall per sign)")
text(s, 7.1, 3.0, 5.6, 3.5, [f"{pct(r)}: " + ", ".join(names) for r, names in sorted(groups.items(), reverse=True)],
     size=15, bullets=True, after=10)

# Backup A2 — reproducibility and AI use
s = new_slide("Backup: Reproducibility and AI Use", "For questions only")
text(s, 0.9, 2.45, 11.6, 4.3, [
    "Code, data and results: github.com/rinrin168/khmersignlanguage",
    "Seed 42 (Python, NumPy, PyTorch), fixed split from src/preprocess.py, torch.save checkpoints",
    "MediaPipe Holistic (Google, Apache 2.0): fixed feature extractor, not trained by me",
    "AI assistance: Claude (Anthropic) for drafting and debugging code, figures and the README; "
    "code and analysis checked by me  [edit this line so it is true for you]",
], size=16, bullets=True, after=12)

prs.save(OUT)
print(f"Saved {OUT}  ({len(prs.slides)} slides; tuned LSTM row included: {tuned})")
